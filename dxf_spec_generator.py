# -*- coding: utf-8 -*-
"""
Генератор спецификаций (xlsx + csv) из папки с DXF-файлами.

Для каждого DXF в выбранной папке формируется строка таблицы:
    Исполнение    — имя файла без расширения
    Материал      — вводится пользователем один раз (общий для всех)
    Расход, м²    — площадь габарита развёртки (ширина × высота), без поворота
    Длина реза, м — суммарная длина всех режущих контуров

Запуск:  python dxf_spec_generator.py
Зависимости: ezdxf, openpyxl (устанавливаются: pip install ezdxf openpyxl)
"""

import csv
import math
import os
import glob
import threading
import traceback

import ezdxf
from ezdxf import bbox

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import openpyxl
from openpyxl.utils import get_column_letter


HEADERS = ["Исполнение", "Материал", "Расход, м²", "Длина реза, м"]

# Коэффициенты перевода единиц DXF ($INSUNITS) в миллиметры.
# 1 = дюймы, 4 = мм, 5 = см, 6 = метры. Остальное считаем миллиметрами.
_UNIT_TO_MM = {1: 25.4, 4: 1.0, 5: 10.0, 6: 1000.0}


def collect_dxf(folder):
    """Отсортированный список путей к .dxf в папке (без учёта регистра)."""
    files = [
        p for p in glob.glob(os.path.join(folder, "*"))
        if p.lower().endswith(".dxf")
    ]
    return sorted(files, key=lambda p: os.path.basename(p).lower())


def _unit_scale(doc):
    """Множитель перевода единиц чертежа в миллиметры."""
    try:
        code = int(doc.header.get("$INSUNITS", 4))
    except (TypeError, ValueError):
        code = 4
    return _UNIT_TO_MM.get(code, 1.0)


def _entity_length(e):
    """Длина одной сущности в единицах чертежа (0, если тип не учитываем)."""
    t = e.dxftype()
    if t == "LINE":
        return e.dxf.start.distance(e.dxf.end)
    if t == "ARC":
        sweep = (e.dxf.end_angle - e.dxf.start_angle) % 360.0
        return math.radians(sweep) * e.dxf.radius
    if t == "CIRCLE":
        return 2.0 * math.pi * e.dxf.radius
    if t in ("LWPOLYLINE", "POLYLINE", "ELLIPSE", "SPLINE"):
        # Универсальный путь: аппроксимация точками через flattening.
        try:
            pts = list(e.flattening(0.1))
        except (AttributeError, TypeError):
            return 0.0
        return sum(a.distance(b) for a, b in zip(pts, pts[1:]))
    return 0.0


def measure(path):
    """Вернуть (area_m2, cut_len_m) для одного DXF-файла."""
    doc = ezdxf.readfile(path)
    msp = doc.modelspace()
    scale = _unit_scale(doc)  # в мм

    ext = bbox.extents(msp)
    width_mm = ext.size.x * scale
    height_mm = ext.size.y * scale
    area_m2 = (width_mm * height_mm) / 1e6

    total_mm = sum(_entity_length(e) for e in msp) * scale
    cut_len_m = total_mm / 1000.0

    return round(area_m2, 4), round(cut_len_m, 3)


def build_rows(folder, material, progress=None):
    """
    Собрать строки таблицы по всем DXF папки.
    Возвращает (rows, errors). progress(done, total, name) — опциональный коллбэк.
    """
    files = collect_dxf(folder)
    rows = []
    errors = []
    total = len(files)
    for i, path in enumerate(files, 1):
        name = os.path.splitext(os.path.basename(path))[0]
        try:
            area_m2, cut_len_m = measure(path)
            rows.append([name, material, area_m2, cut_len_m])
        except Exception as exc:  # битый DXF не должен ронять весь процесс
            rows.append([name, material, "", ""])
            errors.append(f"{os.path.basename(path)}: {exc}")
        if progress:
            progress(i, total, name)
    return rows, errors


def write_xlsx(rows, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    # Авто-ширина колонок по содержимому.
    for col_idx in range(1, len(HEADERS) + 1):
        letter = get_column_letter(col_idx)
        max_len = len(str(HEADERS[col_idx - 1]))
        for r in rows:
            max_len = max(max_len, len(str(r[col_idx - 1])))
        ws.column_dimensions[letter].width = max_len + 2
    wb.save(out_path)


def write_csv(rows, out_path):
    # utf-8-sig + разделитель ';' — чтобы Excel на русской локали разложил по колонкам.
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(HEADERS)
        w.writerows(rows)


# --------------------------------------------------------------------------- GUI

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Генератор спецификаций из DXF")
        self.resizable(False, False)
        self._busy = False
        self.last_folder = None

        pad = {"padx": 8, "pady": 6}
        frm = ttk.Frame(self, padding=12)
        frm.grid(row=0, column=0, sticky="nsew")

        # Папка с DXF
        ttk.Label(frm, text="Папка с DXF:").grid(row=0, column=0, sticky="w", **pad)
        self.folder_var = tk.StringVar()
        ttk.Entry(frm, textvariable=self.folder_var, width=48).grid(
            row=0, column=1, sticky="we", **pad)
        ttk.Button(frm, text="Выбрать…", command=self.choose_folder).grid(
            row=0, column=2, **pad)

        # Материал
        ttk.Label(frm, text="Материал:").grid(row=1, column=0, sticky="w", **pad)
        self.material_var = tk.StringVar()
        ttk.Entry(frm, textvariable=self.material_var, width=48).grid(
            row=1, column=1, columnspan=2, sticky="we", **pad)

        # Кнопки
        self.gen_btn = ttk.Button(frm, text="Сгенерировать", command=self.on_generate)
        self.gen_btn.grid(row=2, column=0, columnspan=2, sticky="we", **pad)

        self.open_btn = ttk.Button(frm, text="Открыть папку",
                                   command=self.open_folder, state="disabled")
        self.open_btn.grid(row=2, column=2, sticky="we", **pad)

        # Статус
        self.status_var = tk.StringVar(value="Выберите папку и укажите материал.")
        ttk.Label(frm, textvariable=self.status_var, foreground="#444",
                  wraplength=460, justify="left").grid(
            row=3, column=0, columnspan=3, sticky="we", **pad)

        frm.columnconfigure(1, weight=1)

    def choose_folder(self):
        folder = filedialog.askdirectory(title="Выберите папку с DXF-файлами")
        if folder:
            self.folder_var.set(folder)

    def open_folder(self):
        if self.last_folder and os.path.isdir(self.last_folder):
            os.startfile(self.last_folder)  # Windows: открыть в Проводнике

    def set_status(self, text):
        self.status_var.set(text)
        self.update_idletasks()

    def on_generate(self):
        if self._busy:
            return
        folder = self.folder_var.get().strip()
        material = self.material_var.get().strip()

        if not folder or not os.path.isdir(folder):
            messagebox.showwarning("Нет папки", "Выберите папку с DXF-файлами.")
            return
        if not material:
            messagebox.showwarning("Нет материала", "Укажите материал.")
            return
        if not collect_dxf(folder):
            messagebox.showwarning("Пусто", "В папке нет .dxf файлов.")
            return

        base = os.path.basename(os.path.normpath(folder))
        xlsx_path = os.path.join(folder, base + ".xlsx")
        csv_path = os.path.join(folder, base + ".csv")

        existing = [p for p in (xlsx_path, csv_path) if os.path.exists(p)]
        if existing:
            names = "\n".join(os.path.basename(p) for p in existing)
            if not messagebox.askyesno("Перезаписать?",
                                       "Уже существуют:\n" + names + "\n\nПерезаписать?"):
                return

        self._busy = True
        self.gen_btn.config(state="disabled")
        threading.Thread(
            target=self._worker,
            args=(folder, material, xlsx_path, csv_path),
            daemon=True,
        ).start()

    def _worker(self, folder, material, xlsx_path, csv_path):
        try:
            def progress(done, total, name):
                self.after(0, self.set_status, f"Обработка {done}/{total}: {name}")

            rows, errors = build_rows(folder, material, progress)
            self.after(0, self.set_status, "Запись файлов…")
            write_xlsx(rows, xlsx_path)
            write_csv(rows, csv_path)

            msg = f"Готово: {len(rows)} строк.\n{xlsx_path}\n{csv_path}"
            if errors:
                msg += "\n\nОшибки чтения:\n" + "\n".join(errors)
            self.last_folder = folder
            self.after(0, self._done, msg, bool(errors))
        except Exception:
            self.after(0, self._fail, traceback.format_exc())

    def _done(self, msg, had_errors):
        self._busy = False
        self.gen_btn.config(state="normal")
        self.open_btn.config(state="normal")
        self.set_status(msg)
        if had_errors:
            messagebox.showwarning("Готово с предупреждениями", msg)
        else:
            messagebox.showinfo("Готово", msg)

    def _fail(self, tb):
        self._busy = False
        self.gen_btn.config(state="normal")
        self.set_status("Ошибка.")
        messagebox.showerror("Ошибка", tb)


if __name__ == "__main__":
    App().mainloop()
